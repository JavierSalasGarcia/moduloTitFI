import os
import sys
import glob
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

RGB_OBSCURO = RGBColor(56, 64, 51) # #384033

INTRODUCCIONES = {
    "Estudiantado con oportunidades para su pleno desarrollo": (
        "El Proyecto 2 consolida a la Facultad de Ingeniería como un referente de excelencia en la formación integral de "
        "profesionales en las diversas ramas de la ingeniería. Su propósito central es garantizar condiciones institucionales "
        "óptimas que favorezcan la permanencia escolar, el egreso oportuno y la titulación, mediante planes de estudio acreditados, "
        "pertinentes y actualizados bajo los criterios transversales del Sistema de Evaluación y Acreditación de la Educación Superior (SEAES). "
        "A través de un sólido programa de tutorías, asesorías disciplinares, infraestructura moderna de laboratorios y oportunidades de movilidad "
        "e internacionalización, se busca formar ingenieras e ingenieros con un alto dominio técnico, pensamiento crítico y una profunda conciencia "
        "social y ética orientada a resolver las grandes problemáticas del entorno."
    ),
    "Estudios avanzados de excelencia": (
        "El Proyecto 3 orienta el desarrollo de los programas de posgrado y especialización hacia la generación de soluciones científicas y "
        "tecnológicas de alto impacto social y productivo. Con un claustro académico integrado mayoritariamente por miembros del Sistema Nacional "
        "de Investigadoras e Investigadores (SNII) y programas registrados en el Sistema Nacional de Posgrados (SNP), la Facultad de Ingeniería "
        "impulsa la formación de recursos humanos de alto nivel en áreas estratégicas como ciencias del agua, ciencia de materiales y ciencias de "
        "la computación. Este proyecto fortalece la gratuidad, las estancias de investigación nacional e internacional, la interdisciplina y la retribución "
        "social, asegurando que las tesis de maestría y doctorado contribuyan directamente al bienestar de las comunidades y a la soberanía científica de la nación."
    ),
    "Investigación abierta, inclusiva": (
        "El Proyecto 4 impulsa la investigación científica, básica y aplicada, como un motor de transformación social y desarrollo tecnológico sustentable. "
        "Centrado en el fortalecimiento de los Cuerpos Académicos, la captación de financiamiento externo y el acceso universal al conocimiento a través del "
        "modelo de Ciencia Abierta, este proyecto articula las capacidades de las investigadoras e investigadores de la Facultad con las necesidades prioritarias "
        "del Estado de México. Asimismo, promueve la equidad de género en las vocaciones científicas STEM, la ética en la investigación y el trabajo colaborativo "
        "en redes del conocimiento de alcance nacional e internacional."
    ),
    "Innovación, desarrollo tecnológico": (
        "El Proyecto 5 articula la vinculación estratégica entre la Facultad de Ingeniería, el sector productivo y las instancias de gobierno para acelerar la "
        "transferencia de conocimientos, patentes y prototipos tecnológicos. Situada en el corredor industrial Toluca-Lerma, la Facultad fortalece la cultura "
        "de la propiedad industrial, el desarrollo de emprendimientos de base tecnológica (spin-offs) y la prestación de servicios especializados de laboratorio "
        "y consultoría técnica, consolidando un ecosistema de innovación que genera valor económico y social con un estricto enfoque de sostenibilidad ambiental."
    ),
    "Difusión y divulgación de la ciencia": (
        "El Proyecto 6 democratiza el conocimiento generado en la Facultad de Ingeniería mediante la divulgación científica rigurosa y el fomento de publicaciones "
        "en revistas de acceso abierto de alto impacto. A través de eventos académicos, coloquios, talleres comunitarios, medios digitales interactivos y "
        "repositorios institucionales abiertos, se busca acercar la ciencia y la tecnología a la sociedad, inspirando vocaciones científicas tempranas en la niñez "
        "y juventud mexiquense y posicionando los bienes públicos del conocimiento al servicio del bien común."
    ),
    "Fortalecimiento de la formación integral del docente": (
        "El Proyecto 9 sitúa al personal académico como el pilar fundamental del quehacer universitario. Su objetivo es asegurar el desarrollo continuo de las "
        "competencias pedagógicas, disciplinares y tecnológicas del profesorado, promoviendo metodologías de enseñanza activas, el uso ético de la inteligencia "
        "artificial y el fortalecimiento del perfil deseable PRODEP y la membresía en el SNII. Asimismo, atiende de manera prioritaria al personal de asignatura "
        "e impulsa esquemas de relevo generacional ordenado, garantizando condiciones de dignificación y excelencia en la práctica docente."
    ),
    "Dignificación de las y los trabajadores": (
        "El Proyecto 10 reconoce la labor indispensable del personal administrativo y operativo en el funcionamiento cotidiano y la calidad de los servicios de la "
        "Facultad. Mediante programas integrales de capacitación laboral, salud ocupacional, seguridad e higiene en laboratorios, renovación ergonómica de espacios "
        "y esquemas transparentes de desarrollo y estímulos, este proyecto promueve un entorno de trabajo digno, seguro y motivante que fortalece el sentido de "
        "pertenencia y el bienestar integral de la comunidad trabajadora."
    ),
    "Gestión administrativa para responder": (
        "El Proyecto 11 moderniza y optimiza los procesos de gestión administrativa y escolar de la Facultad bajo el marco del Sistema de Gestión de la Calidad "
        "ISO 9001:2015. Mediante la simplificación de trámites, la erradicación del uso de papel, la implementación de la firma electrónica avanzada y la "
        "automatización digital de servicios, se garantiza una atención oportuna, transparente, eficiente y centrada en las personas usuarias, fortaleciendo "
        "la calidad de los servicios de apoyo a las funciones sustantivas."
    ),
    "Fortalecimiento tecnológico": (
        "El Proyecto 12 garantiza la vanguardia en infraestructura digital, conectividad, sistemas de información y equipamiento de cómputo en la Facultad de "
        "Ingeniería. Su propósito es dotar a las aulas, laboratorios y áreas académicas de conectividad de alta velocidad, clústeres de supercómputo para "
        "inteligencia artificial y simulación avanzada, licenciamiento de software especializado multiusuario y plataformas seguras, fomentando una cultura digital "
        "integral y de ciberseguridad para enfrentar los desafíos de la Industria 4.0."
    ),
    "Gobernanza": (
        "El Proyecto 13 fortalece la toma de decisiones colegiada, democrática y transparente en la Facultad de Ingeniería a través de sus Consejos de Gobierno "
        "y Académico. Asimismo, consolida la cultura de protección civil, seguridad institucional y gestión integral de riesgos en laboratorios pesados e instalaciones, "
        "preservando la memoria histórica, la identidad universitaria y el orgullo verde y oro de la ingeniería mexiquense."
    ),
    "Reconocimiento de la diversidad": (
        "El Proyecto 18 promueve una cultura de respeto irrestricto a los derechos humanos, la equidad de género, la inclusión y la accesibilidad universal en la "
        "Facultad de Ingeniería. A través de la eliminación de barreras físicas y actitudinales, el impulso a la presencia de mujeres en disciplinas STEM, la "
        "transversalización de la perspectiva de cuidados (como lactarios y apoyos a la corresponsabilidad) y la sensibilización contra cualquier forma de "
        "discriminación, este proyecto garantiza un espacio universitario seguro, justo e incluyente para todas las personas."
    ),
    "Construcción de una cultura de paz": (
        "El Proyecto 19 consolida a la Facultad de Ingeniería como una comunidad universitaria cimentada en la cultura de paz, el diálogo, la resolución no "
        "violenta de controversias y la convivencia armónica. Mediante el fortalecimiento de los mecanismos alternativos de solución de conflictos (MASC), la "
        "vinculación estrecha con la Defensoría de los Derechos Universitarios y la aplicación rigurosa de protocolos de prevención y atención a la violencia "
        "de género, se fomenta un ambiente de confianza, respeto mutuo y bienestar psicoemocional en beneficio de toda la comunidad."
    )
}

def update_introductions_in_docx():
    target_docx = r"D:\Proyectos\plataformaFI\Plan\AcadA\Plantilla_Diseño_PD_ES_FI_2026_2030.docx"
    target_g = r"G:\Para EA\Plantilla_Diseño_PD_ES_FI_2026_2030.docx"
    
    doc = docx.Document(target_docx)
    paragraphs = doc.paragraphs

    updated_count = 0
    for i, p in enumerate(paragraphs):
        p_text = p.text.strip()
        if "Breve introducción" in p_text or "introducción del proyecto" in p_text:
            # Look backwards to find the Heading 2
            h2_text = ""
            for j in range(i - 1, max(0, i - 10), -1):
                if paragraphs[j].style.name.startswith("Heading 2") or any(k.lower() in paragraphs[j].text.lower() for k in INTRODUCCIONES.keys()):
                    h2_text = paragraphs[j].text
                    break
            
            # Find matching intro
            matched_key = None
            for key in INTRODUCCIONES.keys():
                if key.lower() in h2_text.lower():
                    matched_key = key
                    break
            
            if matched_key:
                intro_text = INTRODUCCIONES[matched_key]
                p.text = "" # Clear placeholder
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.line_spacing = 1.15
                
                run = p.add_run(intro_text)
                run.font.name = "DM Sans"
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGB_OBSCURO
                updated_count += 1
                print(f"Updated intro for [{matched_key}]")

    print(f"\nTotal introductions updated: {updated_count}/12")
    doc.save(target_docx)
    try:
        doc.save(target_g)
        print("Saved successfully to G: drive!")
    except Exception as e:
        print(f"Could not save to G: drive: {e}")

if __name__ == "__main__":
    update_introductions_in_docx()
