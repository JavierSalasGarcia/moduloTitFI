#!/usr/bin/env python3
"""Actualiza las series 4*.json con los cortes institucionales entregados en datos2026.

El script sólo publica agregados; nunca copia RFC, CURP ni nombres del libro de personal.
"""
from __future__ import annotations

import json
import math
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "Plan"
DATOS = Path(r"C:\Users\edigi\Downloads\datos2026")
FECHA = "2026-10-22"
PERIODO_911 = "2026B"
PERIODO_EGRESO = "2025-2026"


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(text.upper().strip().split())


def number(value: object) -> int:
    if value in (None, "", "--"):
        return 0
    return int(value)


def sheet_rows(book: Path, sheet: str) -> dict[str, list[object]]:
    wb = load_workbook(book, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        result: dict[str, list[object]] = {}
        for row in ws.iter_rows(values_only=True):
            key = norm(row[0] if row else None)
            if key and key not in result:
                result[key] = list(row)
        return result
    finally:
        wb.close()


def load_json(name: str) -> dict:
    return json.loads((PLAN / name).read_text(encoding="utf-8"))


def save_json(name: str, data: dict) -> None:
    (PLAN / name).write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )


def indicator(section: dict, visual: str) -> dict | None:
    return next((item for item in section.get("indicadores", []) if str(item.get("visual")) == visual), None)


def projects(item: dict) -> list[int]:
    for record in reversed(item.get("registros", [])):
        if isinstance(record, dict) and isinstance(record.get("proyectos"), list):
            return record["proyectos"]
    return []


def set_single(section: dict, visual: str, key: str, value: object) -> None:
    item = indicator(section, visual)
    if item is None:
        return
    item["registros"] = [{key: value, "proyectos": projects(item)}]
    item["numero_registros"] = 1


def upsert_records(item: dict, new_records: list[dict], period_key: str, period: object) -> None:
    kept = [
        record for record in item.get("registros", [])
        if not isinstance(record, dict) or record.get(period_key) != period
    ]
    item["registros"] = kept + new_records
    item["numero_registros"] = len(item["registros"])


def section_source(sections: dict[str, dict], rows: dict[str, list[object]]) -> list[tuple[dict, list[object]]]:
    matched: list[tuple[dict, list[object]]] = []
    for name, section in sections.items():
        row = rows.get(norm(name))
        if row:
            matched.append((section, row))
    return matched


def provenance(data: dict, sources: list[dict], period_note: str) -> None:
    data["fecha_extraccion"] = FECHA
    data["actualizacion_2026"] = {
        "fecha": FECHA,
        "periodo": period_note,
        "fuentes": sources,
        "nota": "Se agregaron únicamente cifras agregadas de la Facultad de Ingeniería; no se publicaron datos personales.",
    }


def totals_from_gender(item: dict, period_key: str, period: object, gender_key: str, value_key: str) -> tuple[int, int, int]:
    records = [r for r in item.get("registros", []) if isinstance(r, dict) and r.get(period_key) == period]
    male = sum(number(r.get(value_key)) for r in records if norm(r.get(gender_key)).startswith("HOMBRE"))
    female = sum(number(r.get(value_key)) for r in records if norm(r.get(gender_key)).startswith("MUJER"))
    return male, female, male + female


def prior_total(item: dict, period_key: str, period: object, value_key: str) -> int:
    order: list[object] = []
    sums: dict[object, int] = {}
    for record in item.get("registros", []):
        if not isinstance(record, dict):
            continue
        p = record.get(period_key)
        if p == period:
            continue
        if p not in sums:
            order.append(p)
            sums[p] = 0
        sums[p] += number(record.get(value_key))
    return sums.get(order[-1], 0) if order else 0


def update_4001(rows: dict[str, list[object]]) -> None:
    name = "4001_solicitudes.json"
    data = load_json(name)
    sections = {"Facultad de Ingeniería": data["facultad"], **data.get("carreras", {})}
    for section, row in section_source(sections, rows):
        male, female, total = number(row[1]), number(row[2]), number(row[3])
        series = indicator(section, "002")
        if not series:
            continue
        prev = prior_total(series, "Ciclo escolar", PERIODO_911, "Suma de Total_Solicitudes")
        proy = projects(series)
        upsert_records(series, [
            {"Ciclo escolar": PERIODO_911, "Solicitudes_Sexo": "Hombres", "Suma de Total_Solicitudes": male, "proyectos": proy},
            {"Ciclo escolar": PERIODO_911, "Solicitudes_Sexo": "Mujeres", "Suma de Total_Solicitudes": female, "proyectos": proy},
        ], "Ciclo escolar", PERIODO_911)
        regional = indicator(section, "013")
        if regional:
            upsert_records(regional, [{"Número y nombre de región": "14  Toluca", "Ciclo escolar": PERIODO_911,
                                       "Suma de Total_Solicitudes": total, "proyectos": projects(regional)}],
                           "Ciclo escolar", PERIODO_911)
        set_single(section, "003", "Superavit Mujeres Solicitudes", female - male)
        set_single(section, "005", "% Crecimiento Anual Solicitudes", (total / prev - 1) if prev else None)
        set_single(section, "007", "Solicitudes Ciclo Actual", total)
        set_single(section, "008", "Indice Paridad Genero Solicitudes1", female / male if male else None)
        set_single(section, "011", "Solicitudes Toluca", total)
        set_single(section, "012", "Solicitudes Fuera de Toluca", 0)
    provenance(data, [{"archivo": "Libro1 e911 (1).xlsx", "hoja": "datos preinsc", "corte": "2026B"}], PERIODO_911)
    save_json(name, data)


def update_4003(rows: dict[str, list[object]]) -> None:
    name = "4003_nuevo_ingreso_uaemex.json"
    data = load_json(name)
    sections = {"Facultad de Ingeniería": data["facultad"], **data.get("carreras", {})}
    for section, row in section_source(sections, rows):
        male, female, total = number(row[1]), number(row[2]), number(row[3])
        series = indicator(section, "003")
        if not series:
            continue
        prev = prior_total(series, "Ciclo escolar", PERIODO_911, "Suma de Total_Ingreso")
        proy = projects(series)
        upsert_records(series, [
            {"Ciclo escolar": PERIODO_911, "Género_Ingreso": "Hombres", "Suma de Total_Ingreso": male, "proyectos": proy},
            {"Ciclo escolar": PERIODO_911, "Género_Ingreso": "Mujeres", "Suma de Total_Ingreso": female, "proyectos": proy},
        ], "Ciclo escolar", PERIODO_911)
        regional = indicator(section, "006")
        if regional:
            upsert_records(regional, [{"Número y nombre de región": "14  Toluca", "Ciclo escolar": PERIODO_911,
                                       "Suma de Total_Ingreso": total, "proyectos": projects(regional)}],
                           "Ciclo escolar", PERIODO_911)
        set_single(section, "001", "Ingreso Fuera de Toluca", 0)
        set_single(section, "002", "Ingreso Ciclo Actual", total)
        set_single(section, "008", "Ingreso Toluca", total)
        set_single(section, "010", "Indice Paridad Genero Ingreso", female / male if male else None)
        set_single(section, "011", "Superavit Mujeres Ingreso", female - male)
        set_single(section, "013", "% Crecimiento Anual Ingreso", (total / prev - 1) if prev else None)
    provenance(data, [{"archivo": "Libro1 e911 (1).xlsx", "hoja": "Matricula total", "corte": "2026B"}], PERIODO_911)
    save_json(name, data)


def update_4005(rows: dict[str, list[object]]) -> None:
    name = "4005_matricula_total.json"
    data = load_json(name)
    sections = {"Facultad de Ingeniería": data["facultad"], **data.get("carreras", {})}
    for section, row in section_source(sections, rows):
        male = number(row[1]) + number(row[4])
        female = number(row[2]) + number(row[5])
        total = number(row[7])
        series = indicator(section, "013")
        if not series:
            continue
        prev = prior_total(series, "Ciclo escolar", PERIODO_911, "Matricula_UAEMEX")
        proy = projects(series)
        upsert_records(series, [
            {"Ciclo escolar": PERIODO_911, "Sexo": "Hombre", "Matricula_UAEMEX": male, "proyectos": proy},
            {"Ciclo escolar": PERIODO_911, "Sexo": "Mujer", "Matricula_UAEMEX": female, "proyectos": proy},
        ], "Ciclo escolar", PERIODO_911)
        set_single(section, "001", "Brecha Genero Absoluta1", female - male)
        set_single(section, "002", "Indice Paridad Genero1", female / male if male else None)
        set_single(section, "005", "Matricula Fuera de Toluca1", 0)
        set_single(section, "010", "% Crecimiento Interanual1", (total / prev - 1) if prev else None)
        set_single(section, "011", "Matricula Toluca1", total)
        set_single(section, "012", "Matricula Ciclo Actual1", total)
    postgraduate_records = [
        ("2018-2019", 103, 41),
        ("2019-2020", 81, 50),
        ("2020-2021", 80, 28),
        ("2021-2022", 67, 31),
        ("2022-2023", 63, 25),
        ("2023-2024", 73, 34),
        ("2024-2025", 72, 27),
        ("2025-2026", 75, 22),
        (PERIODO_911, 70, 11),
    ]
    derived = indicator(data["facultad"], "posgrado_matricula_ingreso")
    if derived is None:
        derived = {
            "visual": "posgrado_matricula_ingreso",
            "metricas": ["Ciclo escolar", "Matrícula total de posgrado", "Nuevo ingreso de posgrado"],
            "campos_modelo": {
                "Serie derivada.Estudios avanzados.Matrícula total": "Matrícula total de posgrado",
                "Serie derivada.Estudios avanzados.Nuevo ingreso": "Nuevo ingreso de posgrado",
            },
            "formatos": {
                "Serie derivada.Estudios avanzados.Matrícula total": "#,0",
                "Serie derivada.Estudios avanzados.Nuevo ingreso": "#,0",
            },
            "registros": [],
            "numero_registros": 0,
        }
        data["facultad"]["indicadores"].append(derived)
    derived["registros"] = [
        {"Ciclo escolar": period, "Matrícula total de posgrado": enrollment,
         "Nuevo ingreso de posgrado": intake, "proyectos": [3, 15]}
        for period, enrollment, intake in postgraduate_records
    ]
    derived["numero_registros"] = len(derived["registros"])
    provenance(data, [{"archivo": "Libro1 e911 (1).xlsx", "hoja": "Matricula total", "corte": "2026B"}], PERIODO_911)
    data["actualizacion_2026"]["serie_posgrado"] = {
        "descripcion": "Agregado de Estudios Avanzados: maestrías y doctorados.",
        "fuentes": ["4003_nuevo_ingreso_uaemex.json", "4005_matricula_total.json",
                    "Libro1 e911 (1).xlsx, hoja Matricula total"],
    }
    save_json(name, data)


def update_4007(rows: dict[str, list[object]]) -> None:
    name = "4007_egresados.json"
    data = load_json(name)
    sections = {"Facultad de Ingeniería": data["facultad"], **data.get("carreras", {})}
    for section, row in section_source(sections, rows):
        male, female, total = number(row[1]), number(row[2]), number(row[3])
        series = indicator(section, "011")
        if not series:
            continue
        prev = prior_total(series, "Ciclo escolar", PERIODO_EGRESO, "Suma de Total_Egresados")
        proy = projects(series)
        upsert_records(series, [
            {"Ciclo escolar": PERIODO_EGRESO, "Sexo": "Hombre", "Suma de Total_Egresados": male, "proyectos": proy},
            {"Ciclo escolar": PERIODO_EGRESO, "Sexo": "Mujer", "Suma de Total_Egresados": female, "proyectos": proy},
        ], "Ciclo escolar", PERIODO_EGRESO)
        set_single(section, "001", "Egresados Ciclo Actual1", total)
        set_single(section, "003", "Indice Paridad Genero Egresados", female / male if male else None)
        set_single(section, "004", "Egresados Fuera de Toluca1", 0)
        set_single(section, "008", "Egresados Toluca1", total)
        set_single(section, "009", "Brecha Genero Absoluta Egresados1", female - male)
        set_single(section, "013", "% Crecimiento Anual Egresados", (total / prev - 1) if prev else None)
    provenance(data, [{"archivo": "Libro1 e911 (1).xlsx", "hoja": "egreso global",
                       "corte": "05/08/2025 al 03/08/2026"}], PERIODO_EGRESO)
    save_json(name, data)


def total_series(section: dict, visual: str, period_key: str, value_key: str) -> list[tuple[object, int]]:
    item = indicator(section, visual)
    if not item:
        return []
    order: list[object] = []
    totals: dict[object, int] = {}
    for record in item.get("registros", []):
        if not isinstance(record, dict) or period_key not in record:
            continue
        p = record[period_key]
        if p not in totals:
            order.append(p)
            totals[p] = 0
        totals[p] += number(record.get(value_key))
    return [(p, totals[p]) for p in order]


def growth_extremes(points: list[tuple[object, int]]) -> tuple[object | None, object | None]:
    growth = [(points[i][0], points[i][1] / points[i - 1][1] - 1)
              for i in range(1, len(points)) if points[i - 1][1]]
    return (max(growth, key=lambda x: x[1])[0], min(growth, key=lambda x: x[1])[0]) if growth else (None, None)


def update_growth_file(source_name: str, target_name: str, source_visual: str, source_period: str,
                       source_value: str, target_visual: str, target_value: str, target_previous: str,
                       period: object, single_map: dict[str, tuple[str, str]]) -> None:
    source = load_json(source_name)
    target = load_json(target_name)
    source_sections = {"Facultad de Ingeniería": source["facultad"], **source.get("carreras", {})}
    target_sections = {"Facultad de Ingeniería": target["facultad"], **target.get("carreras", {})}
    for name, target_section in target_sections.items():
        src_section = source_sections.get(name)
        if not src_section:
            continue
        points = total_series(src_section, source_visual, source_period, source_value)
        if not points or points[-1][0] != period:
            continue
        current = points[-1][1]
        previous = points[-2][1] if len(points) > 1 else 0
        series = indicator(target_section, target_visual)
        if series:
            upsert_records(series, [{source_period: period, target_value: current,
                                     target_previous: previous, "proyectos": projects(series)}], source_period, period)
        initial = points[0][1]
        intervals = len(points) - 1
        maximum, minimum = growth_extremes(points)
        values = {
            "accumulated": current / initial - 1 if initial else None,
            "cagr": (current / initial) ** (1 / intervals) - 1 if initial and intervals else None,
            "maximum": maximum,
            "minimum": minimum,
        }
        for semantic, (visual, key) in single_map.items():
            set_single(target_section, visual, key, values[semantic])
    provenance(target, [{"archivo": source_name, "transformacion": "Tasas recalculadas con el nuevo corte"}], str(period))
    save_json(target_name, target)


def update_4006() -> None:
    source = load_json("4005_matricula_total.json")
    data = load_json("4006_tasa_de_crecimiento_de_matricula.json")
    points = total_series(source["facultad"], "013", "Ciclo escolar", "Matricula_UAEMEX")
    assert points and points[-1][0] == PERIODO_911
    current = points[-1][1]
    previous = points[-2][1]
    series = indicator(data["facultad"], "serie_derivada_de_4005")
    summary = indicator(data["facultad"], "resumen_derivado_de_4005")
    assert series is not None and summary is not None
    upsert_records(series, [{
        "ciclo_escolar": PERIODO_911,
        "matricula": float(current),
        "matricula_anterior": float(previous),
        "crecimiento_interanual_pct": round((current / previous - 1) * 100, 6),
        "proyectos": projects(series),
    }], "ciclo_escolar", PERIODO_911)
    initial = points[0][1]
    intervals = len(points) - 1
    summary["registros"] = [{
        "periodo_inicial": str(points[0][0]),
        "periodo_final": PERIODO_911,
        "crecimiento_acumulado_pct": round((current / initial - 1) * 100, 6),
        "tcma_pct": round(((current / initial) ** (1 / intervals) - 1) * 100, 6),
        "proyectos": projects(summary),
    }]
    summary["numero_registros"] = 1
    provenance(data, [{"archivo": "4005_matricula_total.json",
                       "transformacion": "Tasas recalculadas con el nuevo corte"}], PERIODO_911)
    save_json("4006_tasa_de_crecimiento_de_matricula.json", data)


def update_4031_4034() -> None:
    sources = [{"archivo": "FACULTAD DE INGENIERÍA (2).xls", "hojas": ["BD", "DocentexTiempo"],
                "corte": "Estadística de personal 2026"}]
    academic = load_json("4031_personal_academico.json")
    item = indicator(academic["facultad"], "003")
    assert item is not None
    upsert_records(item, [{
        "Año": 2026,
        "Profesor Tiempo Completo": 74,
        "Profesor Asignatura": 224,
        "Profesor Medio Tiempo": 8,
        "Técnico Académico Medio Tiempo": 0,
        "Técnico Académico Tiempo Completo": 11,
        "Total": 317,
        "proyectos": projects(item),
    }], "Año", 2026)
    provenance(academic, sources, "2026")
    save_json("4031_personal_academico.json", academic)

    growth = load_json("4032_tasa_de_crecimiento_de_pa.json")
    ptc = indicator(growth["facultad"], "001")
    total = indicator(growth["facultad"], "002")
    assert ptc is not None and total is not None
    upsert_records(ptc, [{"Año": 2026, "Total_PTC": 74, "Año Anterior_PTC": 74,
                          "proyectos": projects(ptc)}], "Año", 2026)
    upsert_records(total, [{"Año": 2026, "Total_P Académico": 317,
                            "Año Anterior_Per_Académico": 284, "proyectos": projects(total)}], "Año", 2026)
    provenance(growth, sources, "2026")
    save_json("4032_tasa_de_crecimiento_de_pa.json", growth)

    administrative = load_json("4033_personal_administrativo.json")
    admin = indicator(administrative["facultad"], "003")
    assert admin is not None
    upsert_records(admin, [{"Año": 2026, "Directivo": 1, "Total": 74,
                            "proyectos": projects(admin)}], "Año", 2026)
    provenance(administrative, sources, "2026")
    administrative["actualizacion_2026"]["nota_clasificacion"] = (
        "El libro permite confirmar 73 plazas administrativas y 1 directiva, pero no identifica de forma "
        "equivalente las categorías Confianza y Sindicalizado; por ello sólo se actualizan Directivo y Total."
    )
    save_json("4033_personal_administrativo.json", administrative)

    admin_growth = load_json("4034_tasa_de_crecimiento_de_per_admv.json")
    admin_series = indicator(admin_growth["facultad"], "002")
    assert admin_series is not None
    upsert_records(admin_series, [{"Año": 2026, "Total_P Advo": 74,
                                   "Año Anterior_Per Administrativo": 79,
                                   "proyectos": projects(admin_series)}], "Año", 2026)
    provenance(admin_growth, sources, "2026")
    save_json("4034_tasa_de_crecimiento_de_per_admv.json", admin_growth)


def main() -> None:
    book = DATOS / "Libro1 e911 (1).xlsx"
    applications = sheet_rows(book, "datos preinsc")
    enrollment = sheet_rows(book, "Matricula total")
    graduates = sheet_rows(book, "egreso global")
    update_4001(applications)
    update_4003(enrollment)
    update_4005(enrollment)
    update_4007(graduates)
    update_growth_file(
        "4001_solicitudes.json", "4002_tasa_de_crecimiento_de_solicitudes.json",
        "002", "Ciclo escolar", "Suma de Total_Solicitudes", "008",
        "Solicitudes_UAEMEX", "Ciclo Escolar Anterior_Solicitudes", PERIODO_911,
        {"accumulated": ("001", "Crecimiento Acumulado Total Solicitudes"),
         "cagr": ("003", "TCMA Solicitudes"),
         "maximum": ("009", "Ciclo Mayor Crecimiento Solicitudes"),
         "minimum": ("010", "Ciclo Menor Crecimiento Solicitudes")},
    )
    update_growth_file(
        "4003_nuevo_ingreso_uaemex.json", "4004_tasa_de_crecimiento_de_ingreso.json",
        "003", "Ciclo escolar", "Suma de Total_Ingreso", "010",
        "Ingreso", "Ciclo Escolar Anterior_Ingreso_1", PERIODO_911,
        {"accumulated": ("001", "Crecimiento Acumulado Total Ingreso"),
         "cagr": ("004", "TCMA Ingreso"),
         "maximum": ("003", "Ciclo Mayor Crecimiento Ingreso"),
         "minimum": ("009", "Ciclo Menor Crecimiento Ingreso")},
    )
    update_4006()
    update_growth_file(
        "4007_egresados.json", "4008_tasa_de_crecimiento_egresados.json",
        "011", "Ciclo escolar", "Suma de Total_Egresados", "010",
        "Egresados", "Ciclo Escolar Anterior_Egresados", PERIODO_EGRESO,
        {"accumulated": ("008", "Crecimiento Acumulado Total Egresados"),
         "cagr": ("006", "TCMA Egresados1"),
         "maximum": ("004", "Ciclo Mayor Crecimiento Egresados1"),
         "minimum": ("001", "Ciclo Menor Crecimiento Egresados1")},
    )
    update_4031_4034()
    print("Actualizados: 4001-4008 y 4031-4034 (12 archivos).")


if __name__ == "__main__":
    main()
