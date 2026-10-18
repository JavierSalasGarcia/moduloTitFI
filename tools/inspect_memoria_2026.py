"""Inspección reproducible de la memoria de cálculo; no modifica el XLSX."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

SOURCE = Path(r"C:\Users\edigi\Downloads\Memoria de cálculo NS 2026-2030 Ver4.xlsx")


def value(value):
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def main() -> None:
    formulas = openpyxl.load_workbook(SOURCE, data_only=False, read_only=True)
    cached = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    summary = []
    for name in formulas.sheetnames:
        sheet = formulas[name]
        hit_rows = []
        for row_number, row in enumerate(sheet.iter_rows(), 1):
            row_values = [value(cell.value) for cell in row]
            joined = " | ".join(row_values)
            if re.search(r"Ref\.?\s*2026|2026", joined, re.I):
                hit_rows.append({"row": row_number, "values": row_values})
        summary.append({"sheet": name, "rows": sheet.max_row, "columns": sheet.max_column,
                        "reference_rows": hit_rows[:50], "reference_count": len(hit_rows)})
    out = Path(__file__).resolve().parents[1] / "tmp" / "memoria_2026_inventory.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(summary)} hojas; inventario: {out}")
    print(f"Vínculos externos conservados: {len(getattr(formulas, '_external_links', []))}")
    for item in summary:
        print(f"{item['sheet']!r}: {item['rows']}x{item['columns']}; filas 2026={item['reference_count']}")
    print('\nValores base Ref. 2026 distintos de vacío/cero:')
    for name in cached.sheetnames:
        if not re.match(r'^\d+\s', name):
            continue
        sheet = cached[name]
        for row_number in range(3, sheet.max_row + 1):
            key, label, kind = sheet.cell(row_number, 8).value, sheet.cell(row_number, 12).value, sheet.cell(row_number, 10).value
            raw_base, indicator_base = sheet.cell(row_number, 13).value, sheet.cell(row_number, 18).value
            if raw_base not in (None, '', 0, '-') or indicator_base not in (None, '', 0, '-'):
                print(name, row_number, key, kind, value(label)[:90], 'M=', raw_base, 'R=', indicator_base)
    if len(sys.argv) > 1:
        wanted = sys.argv[1]
        for name in cached.sheetnames:
            if wanted.lower() not in name.lower():
                continue
            print(f"\n### {name}")
            for row_number, row in enumerate(cached[name].iter_rows(values_only=True), 1):
                cells = [f"{openpyxl.utils.get_column_letter(i + 1)}={value(v)}"
                         for i, v in enumerate(row) if v is not None]
                if cells:
                    print(row_number, " | ".join(cells))
            external_formulas = []
            for row in formulas[name].iter_rows():
                for cell in row:
                    raw = str(cell.value or '')
                    if cell.data_type == 'f' and ('[' in raw or '#REF!' in raw):
                        external_formulas.append((getattr(cell, 'coordinate', '?'), raw))
            if external_formulas:
                print('Fórmulas con vínculo o REF:', external_formulas[:30])


if __name__ == "__main__":
    main()
