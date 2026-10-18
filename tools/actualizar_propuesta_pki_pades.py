from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\edigi\Downloads\Propuesta_Firma_Electronica_FI_revisada.docx")
OUTPUT = Path(r"C:\Users\edigi\Downloads\Propuesta_Firma_Electronica_FI_PKI_PAdES.docx")


def find(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"No se encontró el párrafo: {prefix}")


def replace(document: Document, prefix: str, text: str) -> Paragraph:
    paragraph = find(document, prefix)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            paragraph._p.remove(run._r)
    else:
        paragraph.add_run(text)
    return paragraph


def replace_plain(document: Document, prefix: str, text: str) -> Paragraph:
    """Replace all paragraph content, including Word hyperlink elements."""
    paragraph = find(document, prefix)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)
    return paragraph


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


document = Document(SOURCE)

replace(
    document,
    "Se propone a los HH. Consejos Académico",
    "Se propone a los HH. Consejos Académico y de Gobierno de la Facultad de Ingeniería "
    "autorizar el desarrollo y uso gradual de un sistema de firma electrónica integrado a Traza FI, "
    "basado en certificados digitales individuales y firma remota PAdES con una infraestructura de "
    "llave pública interna de la Facultad. Cada firma vinculará matemáticamente el PDF, la identidad "
    "institucional del firmante y su certificado X.509. La llave privada permanecerá cifrada dentro "
    "del servicio de firma y solo podrá utilizarse después de una autorización de segundo factor "
    "ligada al documento específico.",
)

replace(
    document,
    "El objetivo es reducir el tiempo",
    "El objetivo es reducir el tiempo que hoy se invierte en recabar firmas autógrafas, dar "
    "trazabilidad a cada documento y permitir que Adobe Acrobat, Foxit y otros lectores compatibles "
    "detecten alteraciones posteriores. La validez criptográfica se comprobará mediante la firma "
    "PAdES y la cadena de certificados de la Facultad; para que el lector reconozca esa cadena como "
    "confiable, el certificado público de la autoridad raíz deberá instalarse una sola vez en los "
    "equipos institucionales. La adopción será gradual y los procedimientos en papel permanecerán "
    "vigentes durante la transición.",
)

replace(document, "Descripción técnica y alcance actual", "Arquitectura técnica propuesta")

replace(
    document,
    "La operación actual se basa",
    "Traza FI operará como sistema de orquestación y un servicio criptográfico separado realizará "
    "las firmas. El creador cargará uno o varios PDF, definirá firmantes y elegirá una ruta secuencial "
    "o paralela. Cuando llegue su turno, el firmante verá el folio y la huella SHA-256 del documento "
    "y autorizará la operación mediante un segundo factor. El servicio aplicará una firma PAdES en "
    "una actualización incremental del PDF, de modo que cada firma anterior permanezca verificable. "
    "La sesión, los permisos por rol y el token contra falsificación de solicitudes protegerán el "
    "flujo web, pero no sustituirán la verificación criptográfica del segundo factor.",
)

replace(
    document,
    "La representación visual puede ser",
    "La representación visual podrá incluir el nombre del firmante, su cargo o adscripción, fecha y "
    "hora, folio, número de serie del certificado y un enlace o código QR de verificación. La imagen "
    "o el trazo manuscrito seguirá siendo opcional y únicamente visual. La evidencia principal será "
    "la firma digital PAdES incrustada, que contendrá la firma criptográfica y el certificado del "
    "usuario. Por ello, copiar la apariencia a otro archivo no producirá una firma válida.",
)

replace(
    document,
    "El sistema también puede generar con OpenSSL",
    "La PKI interna tendrá una autoridad certificadora raíz de la Facultad y una autoridad "
    "intermedia. La llave privada raíz se generará y conservará fuera de línea, cifrada y bajo "
    "custodia institucional; en el servidor solo operará la intermedia. El certificado público raíz "
    "se distribuirá en formato CER o CRT para instalarlo como ancla de confianza en Adobe Acrobat, "
    "Foxit y, cuando corresponda, en los almacenes institucionales. La llave privada raíz, los "
    "archivos PFX o P12 y cualquier PEM que contenga una llave privada no serán descargables ni "
    "estarán dentro del directorio público del sitio.",
)

future = replace(
    document,
    "Una evolución a firma digital verificable",
    "Cada usuario contará con un certificado X.509 individual emitido por la autoridad intermedia. "
    "Su llave privada se generará en el servicio de firma y se almacenará cifrada fuera de public, "
    "con permisos mínimos y separación respecto de la aplicación web. La autorización para usarla "
    "será de un solo uso, expirará en pocos minutos y quedará ligada al usuario, folio y hash del PDF. "
    "En la primera etapa se empleará PAdES-B-B; posteriormente podrá añadirse PAdES-B-T mediante un "
    "servicio institucional de sellado de tiempo.",
)

insert_after(
    future,
    "La autoridad intermedia publicará información de revocación mediante una lista CRL y, cuando la "
    "infraestructura lo permita, mediante OCSP. La baja de una persona, el cambio de adscripción, la "
    "pérdida de un segundo factor o la sospecha de compromiso provocarán la suspensión o revocación "
    "del certificado. Los PDF conservarán las firmas y certificados ya incorporados para permitir su "
    "validación histórica conforme a la información disponible.",
)

replace(
    document,
    "El análisis jurídico de una firma basada",
    "La arquitectura propuesta utiliza certificados individuales, firma PAdES y autorización "
    "reforzada para aproximar sus controles técnicos a los requisitos de atribución, integridad y "
    "control de los datos de creación de firma. Su clasificación y efectos jurídicos dependerán del "
    "reconocimiento institucional de la PKI, de los procedimientos de identificación y custodia y de "
    "la intervención de la DTIC y del Comité Técnico. Por ello, el piloto no se presentará como firma "
    "electrónica avanzada hasta que esas instancias aprueben la arquitectura y su operación. El marco "
    "aplicable tiene varios niveles que deben valorarse conjuntamente.",
)

replace(
    document,
    "De lo anterior se sigue la ruta jurídica",
    "De lo anterior se sigue la ruta jurídica de esta propuesta. La Facultad podrá construir y probar "
    "una PKI interna sin contratar inicialmente un certificador comercial, pero deberá documentar la "
    "identificación de titulares, emisión, entrega, uso, renovación, suspensión, revocación, "
    "recuperación y destrucción de certificados y llaves. La autoridad raíz, la autoridad intermedia "
    "y el servicio de firma remota deberán coordinarse con la DTIC y obtener el reconocimiento del "
    "Comité Técnico antes de atribuir a las firmas los efectos del artículo 25 de los Lineamientos.",
)

replace(
    document,
    "En la versión actual no existe una llave privada",
    "La administración del servidor no deberá permitir firmar por otra persona. Las llaves privadas "
    "individuales se almacenarán cifradas en un servicio de firma separado del código web y solo se "
    "habilitarán para una operación concreta después de verificar TOTP o, preferentemente, WebAuthn. "
    "La autorización incluirá el hash del PDF y se consumirá una sola vez. La llave raíz permanecerá "
    "fuera de línea; la intermedia tendrá acceso restringido y podrá revocarse si ocurre un incidente. "
    "Los respaldos de llaves estarán cifrados y separados de los respaldos ordinarios del sitio.",
)

replace(
    document,
    "Actualmente, Traza FI registra en MySQL",
    "Cada operación quedará registrada en MySQL con el usuario, certificado, número de serie, fecha y "
    "hora, método de segundo factor, hash de entrada, hash del PDF firmado y resultado de la firma. "
    "La bitácora enlazará criptográficamente sus registros para hacer evidentes las alteraciones y se "
    "respaldará por separado. Se conservarán también las versiones incrementales del PDF, la cadena de "
    "certificados y las listas de revocación necesarias para la validación. La política de retención "
    "deberá aprobarse institucionalmente y contemplar al menos el plazo exigido por los Lineamientos.",
)

replace(
    document,
    "La firma está integrada a Traza FI",
    "La firma se integrará a Traza FI (https://fingenieria.mx). La cuenta institucional identificará "
    "al usuario dentro del flujo, mientras que su certificado individual realizará la firma digital. "
    "Antes de firmar, la plataforma mostrará el documento y exigirá un segundo factor configurado. "
    "El usuario podrá consultar la ruta, descargar el PDF con todas las firmas PAdES y obtener la hoja "
    "de respaldo. El certificado raíz público estará disponible para instalarse en Adobe Acrobat o "
    "Foxit, acompañado de su huella SHA-256 publicada por un canal institucional independiente.",
)

replace(
    document,
    "El beneficio más inmediato es eliminar",
    "El beneficio más inmediato será eliminar los recorridos para recabar firmas y permitir la "
    "verificación criptográfica del resultado. Un documento podrá circular mediante una ruta "
    "secuencial o paralela, conservar cada firma como actualización incremental y mostrar quién tiene "
    "pendiente la actuación. El expediente de Traza FI reunirá los PDF, certificados, hashes, "
    "respuestas, rechazos y eventos. Un verificador por folio o QR complementará la validación del "
    "lector PDF, sin sustituirla.",
)

replace(
    document,
    "PRIMERO. Los HH. Consejos Académico",
    "PRIMERO. Los HH. Consejos Académico y de Gobierno aprueban el desarrollo y pilotaje del sistema "
    "de firma electrónica de Traza FI basado en certificados individuales, firma remota PAdES y una "
    "PKI interna de la Facultad, sujeto a validación técnica, jurídica e institucional antes de su uso "
    "exclusivo.",
)

replace(
    document,
    "SEGUNDO. La Subdirección Académica",
    "SEGUNDO. La Subdirección Académica administrará el piloto y solicitará a la Dirección de "
    "Tecnologías de la Información y Comunicaciones de la UAEMéx la revisión de la arquitectura, la "
    "ceremonia de creación de la autoridad raíz, las políticas de certificados y el reconocimiento "
    "correspondiente ante el Comité Técnico de Firma Electrónica.",
)

replace(
    document,
    "TERCERO. El servicio se alojará",
    "TERCERO. La llave de la autoridad raíz permanecerá fuera de línea. El servidor alojará únicamente "
    "la autoridad intermedia y las llaves individuales cifradas dentro de un servicio de firma "
    "separado, fuera del directorio público, con segundo factor por operación, revocación, respaldos "
    "cifrados, bitácora y medidas de continuidad.",
)

# Referencias técnicas nuevas, incorporadas en orden aproximado.
etsti_anchor = find(document, "Espinosa, A., y de la Paz Reyes")
insert_after(
    etsti_anchor,
    "European Telecommunications Standards Institute. (2024). ETSI EN 319 142-1 V1.2.1: "
    "Electronic Signatures and Infrastructures (ESI); PAdES digital signatures; Part 1: Building "
    "blocks and PAdES baseline signatures. https://www.etsi.org/deliver/etsi_EN/319100_319199/"
    "31914201/01.02.01_60/en_31914201v010201p.pdf",
)

owasp_anchor = find(document, "Miranda, J.")
owasp_paragraph = insert_after(
    owasp_anchor,
    "OWASP Foundation. (2026). Key Management Cheat Sheet. "
    "https://cheatsheetseries.owasp.org/cheatsheets/Key_Management_Cheat_Sheet.html",
)
insert_after(
    owasp_paragraph,
    "pyHanko. (2026). pyHanko documentation. https://docs.pyhanko.eu/",
)

# Desarrollar cada sigla en su primera aparición y explicar brevemente su función.
replace(
    document,
    "Documento que se presenta a los HH.",
    "Documento que se presenta a los Honorables (HH.) Consejos Académico y de Gobierno",
)
replace(
    document,
    "Dr. Javier Salas García",
    "Doctor (Dr., grado académico) Javier Salas García",
)
replace(
    document,
    "Se propone a los HH. Consejos",
    "Se propone a los HH. Consejos Académico y de Gobierno de la Facultad de Ingeniería "
    "autorizar el desarrollo y uso gradual de un sistema de firma electrónica integrado a Traza FI, "
    "donde FI identifica a la Facultad de Ingeniería. El sistema utilizará archivos en formato de "
    "documento portátil (PDF, por sus siglas en inglés), certificados digitales individuales y firma "
    "remota conforme a PDF Advanced Electronic Signatures (PAdES), estándar que incrusta firmas "
    "digitales verificables dentro del PDF. La solución operará con una infraestructura de llave "
    "pública (PKI, por sus siglas en inglés), conjunto de autoridades, certificados y procedimientos "
    "que relacionan identidades con llaves criptográficas. Cada firma vinculará matemáticamente el "
    "PDF, la identidad institucional del firmante y su certificado X.509, estándar que asocia una "
    "identidad con una llave pública. La llave privada permanecerá cifrada dentro del servicio de "
    "firma y solo podrá utilizarse después de una autorización de segundo factor ligada al documento "
    "específico.",
)
replace(
    document,
    "Traza FI operará como sistema",
    "Traza FI operará como sistema de orquestación y un servicio criptográfico separado realizará "
    "las firmas. El creador cargará uno o varios PDF, definirá firmantes y elegirá una ruta secuencial "
    "o paralela. Cuando llegue su turno, el firmante verá el folio y la huella creada con el Algoritmo "
    "de Hash Seguro de 256 bits (SHA-256), función que permite detectar cualquier cambio en el "
    "documento, y autorizará la operación mediante un segundo factor. El servicio aplicará una firma "
    "PAdES en una actualización incremental del PDF, de modo que cada firma anterior permanezca "
    "verificable. La sesión, los permisos por rol y el token contra falsificación de solicitudes "
    "protegerán el flujo web, pero no sustituirán la verificación criptográfica del segundo factor.",
)
replace(
    document,
    "La representación visual podrá incluir",
    "La representación visual podrá incluir el nombre del firmante, su cargo o adscripción, fecha y "
    "hora, folio, número de serie del certificado y un código de respuesta rápida (QR), imagen que "
    "dirige al verificador del sistema. La imagen o el trazo manuscrito seguirá siendo opcional y "
    "únicamente visual. La evidencia principal será la firma digital PAdES incrustada, que contendrá "
    "la firma criptográfica y el certificado del usuario. Por ello, copiar la apariencia a otro "
    "archivo no producirá una firma válida.",
)
replace(
    document,
    "La PKI interna tendrá",
    "La PKI interna tendrá una autoridad certificadora raíz de la Facultad y una autoridad "
    "intermedia. La llave privada raíz se generará y conservará fuera de línea, cifrada y bajo "
    "custodia institucional; en el servidor solo operará la intermedia. El certificado público raíz "
    "se distribuirá como archivo de certificado con extensión .cer o .crt para instalarlo como ancla "
    "de confianza en Adobe Acrobat, Foxit y, cuando corresponda, en los almacenes institucionales. "
    "La llave privada raíz, los contenedores del estándar Public Key Cryptography Standards número 12 "
    "(PKCS#12), normalmente identificados con las extensiones .pfx o .p12, y cualquier archivo "
    "Privacy Enhanced Mail (PEM), formato textual que contenga una llave privada, no serán "
    "descargables ni estarán dentro del directorio público del sitio.",
)
replace(
    document,
    "Cada usuario contará con un certificado",
    "Cada usuario contará con un certificado X.509 individual emitido por la autoridad intermedia. "
    "Su llave privada se generará en el servicio de firma y se almacenará cifrada fuera de public, "
    "con permisos mínimos y separación respecto de la aplicación web. La autorización para usarla "
    "será de un solo uso, expirará en pocos minutos y quedará ligada al usuario, folio y hash del PDF. "
    "En la primera etapa se empleará el perfil PAdES Baseline B (PAdES-B-B), que protege integridad y "
    "autoría con el certificado del firmante; posteriormente podrá añadirse el perfil PAdES Baseline "
    "con sello de tiempo (PAdES-B-T), que incorpora una referencia temporal criptográfica emitida por "
    "un servicio institucional de sellado de tiempo.",
)
replace(
    document,
    "La autoridad intermedia publicará información",
    "La autoridad intermedia publicará información de revocación mediante una lista de certificados "
    "revocados (CRL, por sus siglas en inglés), archivo firmado que enumera certificados que ya no "
    "deben aceptarse, y, cuando la infraestructura lo permita, mediante el Protocolo de Estado de "
    "Certificados en Línea (OCSP, por sus siglas en inglés), servicio que responde el estado de un "
    "certificado en tiempo real. La baja de una persona, el cambio de adscripción, la pérdida de un "
    "segundo factor o la sospecha de compromiso provocarán la suspensión o revocación del certificado. "
    "Los PDF conservarán las firmas y certificados ya incorporados para permitir su validación "
    "histórica conforme a la información disponible.",
)
replace(
    document,
    "La arquitectura propuesta utiliza certificados",
    "La arquitectura propuesta utiliza certificados individuales, firma PAdES y autorización "
    "reforzada para aproximar sus controles técnicos a los requisitos de atribución, integridad y "
    "control de los datos de creación de firma. Su clasificación y efectos jurídicos dependerán del "
    "reconocimiento institucional de la PKI, de los procedimientos de identificación y custodia y de "
    "la intervención de la Dirección de Tecnologías de la Información y Comunicaciones (DTIC), "
    "dependencia universitaria responsable de coordinar la infraestructura tecnológica, y del Comité "
    "Técnico. Por ello, el piloto no se presentará como firma electrónica avanzada hasta que esas "
    "instancias aprueben la arquitectura y su operación. El marco aplicable tiene varios niveles que "
    "deben valorarse conjuntamente.",
)
replace(
    document,
    "En el ámbito federal, el Código de Comercio",
    "En el ámbito federal, el Código de Comercio establece en su artículo 96 que sus disposiciones "
    "no excluirán, restringirán ni privarán de efecto jurídico a ningún método para crear una firma "
    "electrónica, y en su artículo 97 define la firma avanzada o fiable por sus requisitos: que los "
    "datos de creación correspondan exclusivamente al firmante, que estén bajo su control exclusivo "
    "al momento de firmar y que sea posible detectar cualquier alteración de la firma y del documento "
    "(Cámara de Diputados del Honorable [H.] Congreso de la Unión, 2025a). Ninguno de ellos exige un "
    "certificado emitido por un tercero comercial. La propia Secretaría de Economía (2025) aclara "
    "que la firma electrónica simple no requiere certificado digital ni autorización de esa "
    "dependencia, y que la avanzada se distingue por el control exclusivo y la detección de "
    "alteraciones. La Ley de Firma Electrónica Avanzada, publicada en 2012, reconoce que la firma "
    "avanzada produce los mismos efectos que la autógrafa (Cámara de Diputados del H. Congreso de la "
    "Unión, 2025b). El Poder Judicial de la Federación ha sostenido ese criterio en las "
    "jurisprudencias 2a./J. 19/2018 (10a.) y P./J. 5/2021 (11a.), que son sus identificadores "
    "oficiales.",
)
replace(
    document,
    "En el ámbito universitario, la autonomía",
    "En el ámbito universitario, la autonomía que otorga el artículo 3°, fracción VII, de la "
    "Constitución Política de los Estados Unidos Mexicanos (Cámara de Diputados del Honorable [H.] "
    "Congreso de la "
    "Unión, 2026) faculta a la Universidad Autónoma del Estado de México (UAEMéx), institución pública "
    "autónoma a la que pertenece la Facultad, para regular su gestión interna, y la Universidad ya lo "
    "hizo. Los Lineamientos para el Uso de la Firma y Sello Electrónicos de la UAEMéx, vigentes desde "
    "el 16 de octubre de 2012 y publicados en la Gaceta Universitaria número 210 (Universidad Autónoma "
    "del Estado de México, 2012), disponen en su artículo 22 que los documentos con firma electrónica "
    "producen los mismos efectos que los presentados físicamente, y en su artículo 25 que la firma "
    "emitida por el Comité Técnico tiene plena validez para trámites al interior de la Universidad. "
    "El mismo ordenamiento designa a la DTIC como secretaria técnica del Comité y responsable de "
    "emitir los certificados (artículo 14), permite crear comisiones de trabajo en los espacios "
    "académicos (artículo 10) y obliga al titular a mantener el control personal y exclusivo de su "
    "clave privada (artículo 40).",
)
replace(
    document,
    "La UNAM es el antecedente",
    "La Universidad Nacional Autónoma de México (UNAM), institución pública autónoma con una PKI "
    "universitaria propia, es el antecedente más completo en México. Sus lineamientos de 2020 regulan "
    "la Firma Electrónica Universitaria y los servicios de certificación digital que presta su "
    "infraestructura, con una autoridad certificadora institucional que emite los certificados de su "
    "comunidad, idéntica en concepto a la que aquí se propone (UNAM, 2020). Desde 2021, los documentos "
    "firmados con ella producen los mismos efectos que los presentados con firma autógrafa (UNAM, "
    "2021), y desde 2022 se emplea en el sistema de actas de titulación y graduación (UNAM, 2022). Su "
    "Comité de Transparencia adoptó en diciembre de 2021 un Sistema de Firma de Documentos en el que "
    "cada integrante se da de alta como usuario y que prevé volver a la firma autógrafa si el sistema "
    "falla, un esquema de respaldo equivalente al periodo paralelo que se plantea para la Facultad "
    "(Comité de Transparencia de la UNAM, 2021). En el marco de una iniciativa de la Asociación "
    "Nacional de Universidades e Instituciones de Educación Superior (ANUIES), organismo que agrupa a "
    "instituciones mexicanas de educación superior, y la UNAM para digitalizar procesos, la "
    "Universidad de Guadalajara formalizó la operación de una Sub-Autoridad Certificadora (UNAM, "
    "2025).",
)
replace(
    document,
    "La UAEMéx no parte de cero.",
    "La UAEMéx no parte de cero. Su Sistema de Correspondencia Institucional (SICOINS), plataforma "
    "universitaria para gestionar comunicaciones oficiales, ya permite a los titulares firmar "
    "electrónicamente comunicados oficiales (Dirección de Tecnologías de la Información y "
    "Comunicaciones, 2026), y el profesorado firma por medios electrónicos las actas de calificaciones. "
    "La comunidad académica de la Facultad, por tanto, ya firma electrónicamente documentos con plena "
    "validez; lo que falta es extender esa práctica a los oficios, constancias, actas de consejo, "
    "vistos buenos y demás documentos internos que hoy siguen circulando en papel.",
)
replace(
    document,
    "El servicio se alojaría en un servidor",
    "El servicio se alojaría en un servidor contratado con un proveedor externo de nube, fuera de "
    "las instalaciones universitarias, para que la firma siga disponible ante cortes de energía, "
    "intermitencias de red, tomas de instalaciones o daños físicos a equipos. Durante la etapa de "
    "desarrollo el sistema lo administrará la Subdirección Académica de la Facultad, que mediante "
    "oficio coordinará con la DTIC el acceso de esta al servidor, en congruencia con su papel de "
    "secretaria técnica del Comité y emisora de certificados. Se recomienda un proveedor con centro "
    "de datos en México o con cláusulas contractuales de confidencialidad y tratamiento de datos "
    "personales conformes a la ley estatal de la materia, certificación de la Organización "
    "Internacional de Normalización (ISO) 27001, norma para sistemas de gestión de seguridad de la "
    "información, y respaldos automáticos en una segunda región geográfica.",
)
replace(
    document,
    "La administración del servidor no deberá",
    "La administración del servidor no deberá permitir firmar por otra persona. Las llaves privadas "
    "individuales se almacenarán cifradas en un servicio de firma separado del código web y solo se "
    "habilitarán para una operación concreta después de verificar una contraseña de un solo uso basada "
    "en tiempo (TOTP, por sus siglas en inglés), código temporal generado por una aplicación "
    "autenticadora, o, preferentemente, Autenticación Web (WebAuthn), estándar que utiliza una llave "
    "protegida por el dispositivo del usuario. La autorización incluirá el hash del PDF y se consumirá "
    "una sola vez. La llave raíz permanecerá fuera de línea; la intermedia tendrá acceso restringido "
    "y podrá revocarse si ocurre un incidente. Los respaldos de llaves estarán cifrados y separados de "
    "los respaldos ordinarios del sitio.",
)
replace(
    document,
    "Cada operación quedará registrada en MySQL",
    "Cada operación quedará registrada en el sistema gestor de bases de datos MySQL, encargado de "
    "conservar la información estructurada de Traza FI, con el usuario, certificado, número de serie, "
    "fecha y hora, método de segundo factor, hash de entrada, hash del PDF firmado y resultado de la "
    "firma. La bitácora enlazará criptográficamente sus registros para hacer evidentes las "
    "alteraciones y se respaldará por separado. Se conservarán también las versiones incrementales "
    "del PDF, la cadena de certificados y las listas de revocación necesarias para la validación. La "
    "política de retención deberá aprobarse institucionalmente y contemplar al menos el plazo exigido "
    "por los Lineamientos.",
)
replace(
    document,
    "Traza FI es una plataforma web modular",
    "Traza FI es una plataforma web modular de la Facultad de Ingeniería, desarrollada en el lenguaje "
    "de programación del lado del servidor PHP y en MySQL para concentrar servicios de gestión con "
    "acceso según el rol de cada integrante de la comunidad. Integra, entre otros componentes, "
    "solicitudes y tickets, comunicados, encuestas, eventos, archivo documental y rutas de firma. Los "
    "archivos se conservan fuera de la carpeta pública del sitio y se entregan mediante controladores "
    "que verifican la sesión y la participación del usuario. En este contexto, el módulo de firmas "
    "organiza el recorrido de los PDF, permite rutas secuenciales o paralelas y concentra en un "
    "expediente el documento vigente, las respuestas de los firmantes, los rechazos y la hoja de "
    "respaldo.",
)
replace_plain(
    document,
    "Cámara de Diputados del H. Congreso de la Unión. (2025a).",
    "Cámara de Diputados del H. Congreso de la Unión. (2025a). Código de Comercio (última reforma "
    "publicada en el Diario Oficial de la Federación [DOF], medio oficial de publicación normativa, "
    "el 14-11-2025). https://www.diputados.gob.mx/LeyesBiblio/pdf/CCom.pdf",
)
replace_plain(
    document,
    "Espinosa, A., y de la Paz Reyes",
    "Espinosa, A., y de la Paz Reyes, K. (2013, 10 de junio). Firma electrónica prestigia a la "
    "Universidad Veracruzana (UV): comunidad universitaria. UniVerso: El Periódico de los "
    "Universitarios, (525). https://uv.mx/universo-hemeroteca/525/central/central.html",
)
replace_plain(
    document,
    "European Telecommunications Standards Institute.",
    "European Telecommunications Standards Institute (ETSI), organismo europeo que publica normas "
    "técnicas de firma electrónica. (2024). European Standard (EN), denominación de norma europea, "
    "319 142-1 V1.2.1: Electronic Signatures and Infrastructures (ESI), área normativa de firmas e "
    "infraestructuras electrónicas; PAdES digital signatures; Part 1: Building blocks and PAdES baseline "
    "signatures. https://www.etsi.org/deliver/etsi_EN/319100_319199/31914201/01.02.01_60/"
    "en_31914201v010201p.pdf",
)
replace_plain(
    document,
    "OWASP Foundation.",
    "Open Worldwide Application Security Project (OWASP), fundación que publica prácticas abiertas "
    "de seguridad de aplicaciones. (2026). Key Management Cheat Sheet. "
    "https://cheatsheetseries.owasp.org/cheatsheets/Key_Management_Cheat_Sheet.html",
)
replace_plain(
    document,
    "pyHanko. (2026).",
    "pyHanko, biblioteca de software libre para crear y validar firmas digitales en PDF. (2026). "
    "pyHanko documentation. https://docs.pyhanko.eu/",
)
replace_plain(
    document,
    "Universidad Nacional Autónoma de México. (2025).",
    "Universidad Nacional Autónoma de México. (2025). Acta de proyecto: Sub-Autoridad Certificadora "
    "(Sub-AC) UNAM - Universidad de Guadalajara (UdeG), abreviatura institucional de esa universidad. "
    "Dirección General de Cómputo y de Tecnologías de Información y Comunicación. "
    "https://www.fea.unam.mx/wp-content/uploads/2025/02/UNAM-UdeG_ActaProyecto_v3.pdf",
)

replace_plain(
    document,
    "Adobe Community. (s.f.).",
    "Adobe Community. (sin fecha [s.f.], indicación de que la fuente no consigna una fecha de "
    "publicación). How to use digital signature. https://community.adobe.com/t5/acrobat-discussions/"
    "how-to-use-digital-signature/m-p/11643547",
)

for paragraph in document.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True

find(document, "Problemática").paragraph_format.page_break_before = True

document.save(OUTPUT)
print(OUTPUT)
