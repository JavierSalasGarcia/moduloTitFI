"""Extrae la referencia base 2026 sin convertir vacíos en ceros observados."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\edigi\Downloads\Memoria de cálculo NS 2026-2030 Ver4.xlsx")
OUTPUT = ROOT / "Plan" / "memoria_calculo_2026.json"
PROJECTS = list(range(2, 23))


def clean(value):
    return str(value).strip() if value is not None else ""


def project_number(name: str) -> int | None:
    match = re.match(r"^(\d+)\s", name)
    return int(match.group(1)) if match else None


def main() -> None:
    values = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(SOURCE, data_only=False, read_only=True)
    projects = {str(number): {"numero": number, "hoja": None, "metas": [],
                              "referencias_validas": [], "estado": "sin_hoja_en_la_memoria"}
                for number in PROJECTS}
    for sheet_name in values.sheetnames:
        number = project_number(sheet_name)
        if number not in PROJECTS:
            continue
        sheet, formula_sheet = values[sheet_name], formulas[sheet_name]
        project = projects[str(number)]
        project.update({"hoja": sheet_name, "estado": "sin_valor_base_verificable"})
        row = 3
        while row <= sheet.max_row:
            key = clean(sheet.cell(row, 8).value)
            if not key:
                row += 1
                continue
            next_row = row + 1
            while next_row <= sheet.max_row and not clean(sheet.cell(next_row, 8).value):
                next_row += 1
            variable_rows = list(range(row, next_row))
            variables = [{"nombre": clean(sheet.cell(r, 12).value),
                          "valor_ref_2026": sheet.cell(r, 13).value}
                         for r in variable_rows if clean(sheet.cell(r, 12).value)]
            kind = clean(sheet.cell(row, 10).value)
            observed = None
            validation = "no_disponible"
            if kind.lower() == "actividad":
                raw = sheet.cell(row, 13).value
                if isinstance(raw, (int, float)):
                    observed, validation = raw, "valor_directo"
            else:
                supplied = [item["valor_ref_2026"] for item in variables]
                if supplied and all(isinstance(item, (int, float)) for item in supplied):
                    cached = sheet.cell(row, 18).value
                    if isinstance(cached, (int, float)):
                        observed, validation = cached, "calculado_con_variables_completas"
            record = {
                "clave": key, "tipo": kind or "Sin clasificar",
                "meta_prdi": clean(sheet.cell(row, 7).value),
                "meta_espacio_academico": clean(sheet.cell(row, 9).value),
                "metodo_calculo": clean(sheet.cell(row, 11).value),
                "variables": variables, "valor_ref_2026": observed,
                "estado_ref_2026": validation,
                "formula_ref_2026": clean(formula_sheet.cell(row, 18).value)
                    if formula_sheet.cell(row, 18).data_type == "f" else None,
                "proyeccion_libro": {str(year): sheet.cell(row, col).value
                                     for year, col in zip(range(2027, 2031), range(19, 23))},
            }
            project["metas"].append(record)
            if isinstance(observed, (int, float)) and observed > 0:
                project["referencias_validas"].append({"clave": key, "valor": observed,
                                                       "descripcion": variables[0]["nombre"] if variables else record["meta_prdi"]})
            elif observed == 0:
                project.setdefault("ceros_por_validar", []).append({"clave": key,
                    "descripcion": variables[0]["nombre"] if variables else record["meta_prdi"],
                    "motivo": "El cero está escrito en Ref. 2026, pero debe confirmarse con el área responsable antes de tratarlo como observación."})
            row = next_row
        if project["referencias_validas"]:
            project["estado"] = "con_valor_base_verificable"
    total_metas = sum(len(item["metas"]) for item in projects.values())
    total_bases_positivas = sum(len(item["referencias_validas"]) for item in projects.values())
    total_resumen = values["Hoja1"].cell(20, 4).value
    data = {
        "titulo": "Referencia base 2026 para la planeación de la Facultad de Ingeniería",
        "version_archivo": 4,
        "control_totales": {
            "filas_detalladas_extraidas": total_metas,
            "total_declarado_en_hoja_resumen": total_resumen,
            "diferencia": total_resumen - total_metas if isinstance(total_resumen, (int, float)) else None,
        },
        "periodo_base": "Ref. 2026", "autor_corporativo": "Sistema de Planeación y Desarrollo Institucional",
        "sigla": "SPDI", "anio": 2026,
        "cita_apa_primera": "(Sistema de Planeación y Desarrollo Institucional [SPDI], 2026a)",
        "cita_apa_subsecuente": "(SPDI, 2026a)",
        "referencia_apa": "Sistema de Planeación y Desarrollo Institucional. (2026a). Memoria de cálculo NS 2026–2030 (versión 4) [Hoja de cálculo no publicada]. Universidad Autónoma del Estado de México.",
        "archivo_origen": str(SOURCE),
        "criterio_integracion": ("Se adopta Ref. 2026 como referencia base de planeación. Las celdas vacías, fórmulas sin variables completas y errores #REF! "
                                 "se registran como dato no disponible, no como cero. Los valores de años posteriores son programación y no observaciones."),
        "hallazgos_calidad": [
            f"Se extrajeron {total_metas} actividades e indicadores de las hojas de proyecto; {total_bases_positivas} tienen un valor Ref. 2026 positivo y verificable con el criterio aplicado.",
            f"La hoja de resumen declara {total_resumen} metas, pero las hojas detalladas contienen {total_metas}: el proyecto 2 declara 27 y detalla 26; el proyecto 6 declara 4 y detalla 3. La diferencia requiere conciliación con SPDI.",
            "La hoja Indicadores x proyecto contiene referencias #REF! y no se utiliza como fuente numérica.",
            "La mayoría de las metas no tiene variables completas para calcular Ref. 2026.",
            "Sólo se incorporan como observados los valores directos o indicadores con todas sus variables numéricas.",
        ],
        "cambios_detectados_respecto_version_3": [
            "La celda de meta del espacio académico del proyecto 2 incorpora el texto provisional “meta fac”; se conserva como texto de trabajo y no se interpreta como meta aprobada.",
            "La fila ACT-POA-NS-2651 del proyecto 19 cambia su clasificación de indicador a actividad; su Ref. 2026 de cinco se integra como valor directo.",
            "No se detectaron nuevos valores positivos de Ref. 2026 en las celdas de variables o resultados respecto de la versión 3.",
        ],
        "impacto_en_el_diagnostico": {
            "uso": "Ref. 2026 sustituye cualquier referencia anterior como línea base de las metas del plan.",
            "contexto_historico": "Las series de los tableros institucionales anteriores a 2026 se conservan para explicar tendencias, brechas entre licenciaturas y valores extremos, pero no funcionan como línea base de programación.",
            "regla_de_calidad": "Una referencia sólo se incorpora como observación cuando existe un valor directo o cuando todas las variables de la fórmula están disponibles. Los vacíos, errores #REF! y fórmulas incompletas requieren validación del área responsable.",
            "efecto_en_metas": "Las metas cuantitativas se mantienen iguales o por encima de Ref. 2026. Cuando la memoria no contiene una referencia verificable, el primer resultado esperado es documentarla y validarla antes de fijar una cantidad de crecimiento.",
        },
        "proyectos": projects,
    }
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    valid = [(number, item["referencias_validas"]) for number, item in projects.items() if item["referencias_validas"]]
    print(f"Generados {len(projects)} proyectos; referencias base verificables: {valid}")


if __name__ == "__main__":
    main()
