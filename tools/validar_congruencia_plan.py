#!/usr/bin/env python3
"""Audita la congruencia entre diagnóstico, matriz y memoria de cálculo."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from actualizar_matriz_consistencia_desde_sitio import (
    DEFAULT_MATRIX,
    DEFAULT_MEMORY,
    project_sheets,
    validate_matrix,
    validate_memory,
)
from actualizar_memoria_calculo_desde_sitio import load_projects, rows_by_key
from plan_memoria_metodologia import complete_projects, direct_indicator_baselines, historical_context, key


ROOT = Path(__file__).resolve().parents[1]
PLAN_DIR = ROOT / "Plan"


def main() -> int:
    projects, _, _ = load_projects(True, "", None)
    projects, report = complete_projects(projects, PLAN_DIR)

    matrix = load_workbook(DEFAULT_MATRIX, data_only=False)
    validate_matrix(matrix, projects, report)
    validate_memory(DEFAULT_MEMORY, projects, report)

    formulas = load_workbook(DEFAULT_MEMORY, data_only=False, read_only=True)
    values = load_workbook(DEFAULT_MEMORY, data_only=True, read_only=True)
    formula_sheets = project_sheets(formulas)
    value_sheets = project_sheets(values)
    missing = []
    errors = []
    outputs: dict[str, list[float]] = {}

    policy_count = 0
    obligation = re.compile(
        r"\b(?:deberá|deberán|tendrá|tendrán|será|serán|entrará|entrarán|garantizará|garantizarán)\b|"
        r"\bse\s+(?:destinará|destinarán|orientará|orientarán|regirá|regirán)\b",
        flags=re.I,
    )
    policy_subject = re.compile(
        r"^(?:El|La|Los|Las)\s+(?:personal|personas?|estudiantado|profesorado|comunidad|"
        r"comités?|coordinaciones?|comisiones?|núcleos?|órganos?)\b",
        flags=re.I,
    )
    forbidden_policy = re.compile(
        r"\bOE\d+\b|\bobjetivo\b|\bmetas?\s+anuales?\b|\bcada\s+año\b|"
        r"\b20(?:2[6-9]|30)\b|se programarán con responsables|"
        r"tendrán responsable, plazo e indicador|antes de aprobar cambios",
        flags=re.I,
    )
    for project in projects:
        policies = project.get("politicas", [])
        policy_count += len(policies)
        covered = {objective for policy in policies for objective in policy.get("objetivos", [])}
        expected = {objective.get("id") for objective in project.get("objetivos_especificos", [])}
        if covered != expected:
            errors.append(
                f"Proyecto {project['numero']}: políticas asignadas a {sorted(covered)}, se esperaban {sorted(expected)}"
            )
        if len(policies) < 2:
            errors.append(f"Proyecto {project['numero']}: menos de dos políticas")
        for policy in policies:
            text = str(policy.get("texto", "")).strip()
            if not obligation.search(text):
                errors.append(f"Proyecto {project['numero']} {policy.get('id')}: no denota obligatoriedad futura")
            if not policy_subject.search(text):
                errors.append(
                    f"Proyecto {project['numero']} {policy.get('id')}: no identifica la población o instancia "
                    "específica que asume la obligación"
                )
            if forbidden_policy.search(text):
                errors.append(f"Proyecto {project['numero']} {policy.get('id')}: contiene lenguaje operativo o redundante")
            if re.match(r"^(Realizar|Aplicar|Elaborar|Registrar|Actualizar|Revisar|Coordinar|Implementar)\b", text, flags=re.I):
                errors.append(f"Proyecto {project['numero']} {policy.get('id')}: está redactada como estrategia o actividad")

    for project in projects:
        metas = project.get("referencia_base_2026", {}).get("metas", [])
        if not metas:
            continue
        number = int(project["numero"])
        formula_sheet = formula_sheets[number]
        value_sheet = value_sheets[number]
        rows = rows_by_key(formula_sheet)
        for item in metas:
            item_key = key(item.get("clave"))
            row = rows[item_key]
            for offset, _variable in enumerate(item.get("variables", [])):
                for column in range(14, 18):
                    if formula_sheet.cell(row + offset, column).value is None:
                        missing.append(formula_sheet.cell(row + offset, column).coordinate)
            result = []
            for column in range(19, 23):
                formula = formula_sheet.cell(row, column).value
                value = value_sheet.cell(row, column).value
                if not isinstance(formula, str) or not formula.startswith("="):
                    errors.append(f"{formula_sheet.title}!{formula_sheet.cell(row, column).coordinate}: sin fórmula")
                if isinstance(formula, str) and "#REF!" in formula.upper():
                    errors.append(f"{formula_sheet.title}!{formula_sheet.cell(row, column).coordinate}: #REF!")
                if isinstance(value, str) and value.startswith("#"):
                    errors.append(f"{value_sheet.title}!{value_sheet.cell(row, column).coordinate}: {value}")
                if value is None:
                    errors.append(f"{value_sheet.title}!{value_sheet.cell(row, column).coordinate}: sin cálculo")
                else:
                    result.append(float(value))
            outputs[item_key] = result

    context = historical_context(PLAN_DIR)
    baselines = direct_indicator_baselines(PLAN_DIR, context)
    for item_key in ("IND-PRDI-353", "IND-PRDI-40", "IND-NS-127", "IND-PRDI-221"):
        series = outputs[item_key]
        if min(series) + 0.011 < baselines[item_key]:
            errors.append(f"{item_key}: alguna meta es menor que la línea base {baselines[item_key]:.2f}")
        if any(current + 0.011 < previous for previous, current in zip(series, series[1:])):
            errors.append(f"{item_key}: trayectoria no creciente")
    abandonment = outputs["IND-PRDI-34"]
    if max(abandonment) - 0.011 > baselines["IND-PRDI-34"]:
        errors.append("IND-PRDI-34: el abandono supera la línea base histórica")
    if any(current - 0.011 > previous for previous, current in zip(abandonment, abandonment[1:])):
        errors.append("IND-PRDI-34: el abandono no disminuye de forma gradual")
    enrollment = outputs["ACT-PRDI-NS-1764"]
    if min(enrollment) < context["matricula"]:
        errors.append("ACT-PRDI-NS-1764: la matrícula programada es menor que la histórica")

    additional = values["Metas adicionales"]
    additional_projects = {
        int(str(additional.cell(row, 1).value).split(".", 1)[0])
        for row in range(3, additional.max_row + 1) if additional.cell(row, 2).value
    }
    if additional_projects != {14, 15, 17}:
        errors.append(f"Proyectos excepcionales inesperados: {sorted(additional_projects)}")
    additional_count = sum(
        1 for row in range(3, additional.max_row + 1) if additional.cell(row, 2).value
    )
    if additional_count != 9:
        errors.append(f"Se encontraron {additional_count} metas adicionales, no 9")
    if policy_count != 54:
        errors.append(f"Se encontraron {policy_count} políticas, no 54")

    if missing or errors:
        detail = "; ".join((missing + errors)[:30])
        raise RuntimeError(f"La auditoría falló: {detail}")

    print("Auditoría triangular superada.")
    print("21 proyectos y 121 renglones PRDI congruentes con el diagnóstico y la matriz.")
    print("130 metas de matriz: 121 corresponden literalmente a la Memoria y 9 son excepcionales (proyectos 14, 15 y 17).")
    print("54 políticas identifican un sujeto específico, expresan obligación futura y no contienen referencias a objetivos, metas o calendarios.")
    print("Las trayectorias históricas comparables son conservadoras y no retroceden.")
    print(f"Matrícula histórica: {context['matricula']:.0f}; programación: {outputs['ACT-PRDI-NS-1764']}")
    print(f"Cobertura de tutoría base: {baselines['IND-PRDI-353']:.2f}%; programación: {outputs['IND-PRDI-353']}")
    print(f"Abandono base: {baselines['IND-PRDI-34']:.2f}%; programación: {outputs['IND-PRDI-34']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
