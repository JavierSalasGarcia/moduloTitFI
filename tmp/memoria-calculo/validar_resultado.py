from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula


ROOT = Path(r"D:\Proyectos\plataformaFI")
TARGET = ROOT / "Plan" / "Memoria de cálculo NS 2026-2030Ultima.xlsx"
BACKUP = ROOT / "Plan" / "respaldos" / "Memoria de cálculo NS 2026-2030Ultima-20260923-142541.xlsx"


def formula_errors(workbook):
    errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(token in value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                    errors.append((sheet.title, cell.coordinate, value))
    return errors


def comparable(value):
    if isinstance(value, ArrayFormula):
        return ("array-formula", value.ref, value.text)
    return value


before = load_workbook(BACKUP, data_only=False)
after = load_workbook(TARGET, data_only=False)
cached = load_workbook(TARGET, data_only=True)

assert before.sheetnames == after.sheetnames, "Cambió la lista u orden de hojas."
assert len(after.sheetnames) == 22

structural_changes = []
value_changes = {}
for old_sheet, new_sheet in zip(before.worksheets, after.worksheets):
    if set(old_sheet.merged_cells.ranges) != set(new_sheet.merged_cells.ranges):
        structural_changes.append(f"combinaciones:{new_sheet.title}")
    if len(old_sheet._charts) != len(new_sheet._charts):
        structural_changes.append(f"gráficas:{new_sheet.title}")
    if len(old_sheet._images) != len(new_sheet._images):
        structural_changes.append(f"imágenes:{new_sheet.title}")
    changed = []
    max_row = max(old_sheet.max_row, new_sheet.max_row)
    max_col = max(old_sheet.max_column, new_sheet.max_column)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            if comparable(old_sheet.cell(row, col).value) != comparable(new_sheet.cell(row, col).value):
                changed.append(new_sheet.cell(row, col).coordinate)
    if changed:
        value_changes[new_sheet.title] = changed

assert not structural_changes, structural_changes
unexpected = set(value_changes) - {"4 Investigación abierta", "19 Constr cultura de paz ", "Metas adicionales"}
if unexpected:
    for sheet_name in sorted(unexpected):
        print(f"Cambios inesperados en {sheet_name}:")
        for coordinate in value_changes[sheet_name][:20]:
            print(f"  {coordinate}: {before[sheet_name][coordinate].value!r} -> {after[sheet_name][coordinate].value!r}")
    raise AssertionError(f"Hojas con cambios inesperados: {sorted(unexpected)}")

additional = after["Metas adicionales"]
rows = [row for row in range(3, additional.max_row + 1) if additional.cell(row, 2).value]
assert len(rows) == 61
assert len({additional.cell(row, 1).value.split(".", 1)[0] for row in rows}) == 21
for row in rows:
    assert additional.cell(row, 17).value == "Por validar con el área responsable"
    for output_col, input_col in zip(range(11, 16), range(6, 11)):
        letter = additional.cell(row, input_col).column_letter
        assert additional.cell(row, output_col).value == f'=IF({letter}{row}="","",{letter}{row})'

checks = [
    ("4 Investigación abierta", 3, "ACT-PRDI-NS-2870", 3),
    ("19 Constr cultura de paz ", 3, "ACT-POA-NS-2651", 5),
]
for sheet_name, row, key, expected in checks:
    sheet = after[sheet_name]
    assert sheet.cell(row, 8).value == key
    assert sheet.cell(row, 13).value == expected
    assert [sheet.cell(row, col).value for col in range(14, 18)] == [expected] * 4
    assert sheet.cell(row, 24).value == "Por validar con el área responsable"

before_errors = formula_errors(before)
after_errors = formula_errors(after)
assert before_errors == after_errors, "Cambió el conjunto de errores preexistentes en fórmulas."

cached_nonempty = 0
cached_additional = cached["Metas adicionales"]
for row in rows:
    cached_nonempty += sum(cached_additional.cell(row, col).value is not None for col in range(11, 16))

print(f"Hojas conservadas: {len(after.sheetnames)}")
print("Cambios limitados a:")
for sheet_name, cells in value_changes.items():
    print(f"  {sheet_name}: {len(cells)} celdas con valor/fórmula modificado")
print(f"Metas adicionales: {len(rows)} en {len(set(additional.cell(r, 1).value for r in rows))} proyectos")
print("Coincidencias exactas verificadas: 2")
print(f"Errores de fórmula preexistentes sin cambio: {len(after_errors)}")
print(f"Resultados calculados no vacíos en K:O: {cached_nonempty}")
